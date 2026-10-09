#!/usr/bin/env python3
# Чувствительность базового сценария финмодели: python3 docs/business/sensitivity.py
#
# Берёт finmodel.xlsx, меняет по одному допущению на листе «Допущения»,
# пересчитывает копии в LibreOffice и печатает итоги базового сценария за три
# года таблицей Markdown. Итоги — оба сегмента: МСБ и крупные события.
# Числа из этой таблицы должны стоять в docs/business/interviews/assumptions.md
# (обновить после изменения финмодели).
# Сам finmodel.xlsx не меняется.
#
# Нужны openpyxl и LibreOffice (soffice).
import os
import subprocess
import sys
import tempfile

import openpyxl

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'finmodel.xlsx')

# Ячейки листа «Допущения»: колонка D — базовый сценарий, C — общие параметры.
# Строки крупного сегмента ищутся по подписи в колонке B: ключ вида
# ('Крупных событий в год, год 1', 'D') превращается в адрес ячейки.
LEV = ('Крупных событий в год, год 1', 'Крупных событий в год, год 2', 'Крупных событий в год, год 3')
LTPE = 'Продано билетов на крупное событие'
LPRICE = 'Средняя цена билета на крупном событии'
LCOST = 'Продажа и сопровождение одного крупного события'

CASES = [
    ('Базовый сценарий', '—', {}),
    ('Цена билета', '4 000 ₸', {'D22': 4000}),
    ('Цена билета', '8 000 ₸', {'D22': 8000}),
    ('Билетов на событие', '60', {'D29': 60}),
    ('Билетов на событие', '160', {'D29': 160}),
    ('Событий на организатора в месяц', '0,7', {'D28': 0.7}),
    ('Событий на организатора в месяц', '2,0', {'D28': 2.0}),
    ('Отток организаторов в месяц', '8%', {'D27': 0.08}),
    ('Отток организаторов в месяц', '2,5%', {'D27': 0.025}),
    ('Новых организаторов в месяц', 'вдвое меньше', {'D24': 1.5, 'D25': 2.5, 'D26': 3.5}),
    ('Новых организаторов в месяц', 'в 1,5 раза больше', {'D24': 4.5, 'D25': 7.5, 'D26': 10.5}),
    ('CAC', '50 000 ₸', {'D30': 50000}),
    ('CAC', '300 000 ₸', {'D30': 300000}),
    ('Эквайринг', '1,2%', {'D23': 0.012}),
    ('Эквайринг', '3%', {'D23': 0.03}),
    ('Доля бесплатных билетов', '0%', {'C12': 0.0}),
    ('Доля бесплатных билетов', '30%', {'C12': 0.30}),
    ('Доля возвратов', '10%', {'C13': 0.10}),
    ('Билетов в заказе', '1,3', {'C9': 1.3}),
    ('Билетов в заказе', '3', {'C9': 3}),
    ('SMS на заказ', '0 ₸', {'C10': 0}),
    ('SMS на заказ', '40 ₸', {'C10': 40}),
    ('Без крупного сегмента', '0 событий', {(LEV[0], 'D'): 0, (LEV[1], 'D'): 0, (LEV[2], 'D'): 0}),
    ('Крупных событий вдвое больше', '0 / 2 / 4', {(LEV[0], 'D'): 0, (LEV[1], 'D'): 2, (LEV[2], 'D'): 4}),
    ('Билетов на крупное событие', '4 000', {(LTPE, 'D'): 4000}),
    ('Билетов на крупное событие', '12 000', {(LTPE, 'D'): 12000}),
    ('Цена билета на крупном событии', '10 000 ₸', {(LPRICE, 'D'): 10000}),
    ('Цена билета на крупном событии', '25 000 ₸', {(LPRICE, 'D'): 25000}),
    ('Продажа и сопровождение события', '600 000 ₸', {(LCOST, 'D'): 600000}),
    ('Продажа и сопровождение события', '2 500 000 ₸', {(LCOST, 'D'): 2500000}),
]


def cell_address(ws, key):
    # Строка ищется по подписи в колонке B, чтобы не зависеть от номера строки.
    if isinstance(key, str):
        return key
    label, col = key
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 2).value == label:
            return f'{col}{row}'
    raise KeyError(f'нет строки «{label}» на листе «Допущения»')


def mln(v):
    return f'{v / 1e6:+.1f}'.replace('.', ',').replace('-', '−')


def num(v, digits=1):
    return f'{v:.{digits}f}'.replace('.', ',')


def main():
    with tempfile.TemporaryDirectory() as tmp:
        src_dir, out_dir = os.path.join(tmp, 'in'), os.path.join(tmp, 'out')
        os.makedirs(src_dir)
        paths = []
        for i, (_, _, edits) in enumerate(CASES):
            wb = openpyxl.load_workbook(SRC)
            ws = wb['Допущения']
            for key, value in edits.items():
                ws[cell_address(ws, key)] = value
            p = os.path.join(src_dir, f'case{i:02d}.xlsx')
            wb.save(p)
            paths.append(p)
        # openpyxl сохраняет формулы без значений, LibreOffice считает их при открытии.
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out_dir, *paths],
                       check=True, stdout=subprocess.DEVNULL, timeout=600)
        print('| Допущение | Значение | Операционная прибыль за 3 года, млн ₸ | Вложить, млн ₸ | Окупаемость, месяц | LTV/CAC | Маржа на билет МСБ, ₸ |')
        print('|---|---|---:|---:|---:|---:|---:|')
        for i, (label, value, _) in enumerate(CASES):
            wb = openpyxl.load_workbook(os.path.join(out_dir, f'case{i:02d}.xlsx'), data_only=True)
            s, u = wb['Итоги'], wb['Юнит-экономика']
            payback = s['E16'].value
            payback = payback if isinstance(payback, (int, float)) else 'нет за 3 года'
            print(f'| {label} | {value} | {mln(s["E19"].value)} | {num(s["E17"].value / 1e6)} | {payback} '
                  f'| {num(s["E20"].value)} | {num(u["D11"].value, 0)} |')
    return 0


if __name__ == '__main__':
    sys.exit(main())
