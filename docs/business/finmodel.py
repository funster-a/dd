# Финансовая модель «Партера»: 3 года помесячно, три сценария.
# Собирает docs/business/finmodel.xlsx: python3 docs/business/finmodel.py
# Значения формул считает Excel или LibreOffice при открытии.
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.chart import LineChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'finmodel.xlsx')
F = 'Arial'
BLUE = Font(name=F, color='0000FF')
BLACK = Font(name=F, color='000000')
GREEN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
H2 = Font(name=F, bold=True, size=11)
MUTED = Font(name=F, color='666666', italic=True, size=9)
YELLOW = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='EDEBE6')
THIN = Border(bottom=Side(style='thin', color='999999'))
TENGE = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'

wb = Workbook()

def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None,
                              italic=c.font.italic if c.font else False, size=c.font.size if c.font else 11)

# ---------------------------------------------------------------- Допущения
A = wb.active
A.title = 'Допущения'
A['A1'] = 'Финансовая модель «Партер» — допущения'; A['A1'].font = TITLE
A['A2'] = ('Синие числа — вводимые значения, их можно менять; жёлтым выделены ключевые бизнес-параметры. '
           'Все остальные листы считаются формулами от этого листа. Суммы — в тенге, проценты — доли.')
A['A2'].font = MUTED
A.column_dimensions['A'].width = 4
A.column_dimensions['B'].width = 52
for col in 'CDE':
    A.column_dimensions[col].width = 16
A.column_dimensions['F'].width = 12
A.column_dimensions['G'].width = 90

names = {}  # имя → абсолютная ссылка

def inp(row, key, label, value, fmt, unit, note, key_param=False):
    A.cell(row, 2, label).font = BLACK
    c = A.cell(row, 3, value); c.font = BLUE; c.number_format = fmt
    if key_param: c.fill = YELLOW
    A.cell(row, 6, unit).font = MUTED
    A.cell(row, 7, note).font = MUTED
    names[key] = f"Допущения!$C${row}"

r = 4
A.cell(r, 2, 'Общие параметры').font = H2; A.cell(r, 3, 'Значение').font = BOLD; A.cell(r, 7, 'Источник / обоснование').font = BOLD
r += 1
inp(r, 'start', 'Первый месяц модели', date(2027, 6, 1), 'mmm yyyy', '', 'Запуск после защиты диплома (май 2027). Допущение.'); r += 1
inp(r, 'fee', 'Сервисный сбор с покупателя', 0.05, PCT, 'от цены', 'Бизнес-решение 2026-10-02: 5% (сначала 7%, снижено владельцем продукта). Рынок ~10%. В продукте — SERVICE_FEE_BPS, ADR 019.', True); r += 1
inp(r, 'orgfee', 'Комиссия с организатора', 0.0, PCT, 'от цены', 'Бизнес-решение 2026-10-02: организатор не платит.', True); r += 1
inp(r, 'tax', 'Налог (упрощённый режим, ИП)', 0.04, PCT, 'от дохода', 'Базовая ставка 4% по Налоговому кодексу РК с 2026; акимат может снизить до 2%. Доход — только сбор: нужна агентская схема (лист «Риски»).'); r += 1
inp(r, 'tpo', 'Билетов в одном заказе', 2.0, NUM1, 'шт.', 'Допущение: пара — типичный заказ на концерт/стендап.'); r += 1
inp(r, 'sms', 'SMS с кодом входа на заказ', 20, TENGE, '₸', 'Допущение по местным SMS-агрегаторам; международные шлюзы дороже (~€0,14). Код нужен при входе, не на каждый заказ — консервативно.'); r += 1
inp(r, 'infra', 'Инфраструктура и письма на билет', 5, TENGE, '₸', 'Допущение: почта, хранилище, вычисления сверх фиксированного хостинга.'); r += 1
inp(r, 'free', 'Доля бесплатных билетов', 0.10, PCT, 'от проданных', 'Бесплатные билеты есть в продукте; сбор с них не берётся. Допущение.'); r += 1
inp(r, 'refund', 'Доля возвращённых билетов', 0.03, PCT, 'от платных', 'Допущение. С запасом: в продукте сбор возвращается только при отмене события, при возврате по желанию покупателя он остаётся (ADR 019); эквайринг не возвращается.'); r += 1
inp(r, 'legal1', 'Юр. сопровождение запуска (разово, 1-й месяц)', 800000, TENGE, '₸', 'Агентский договор, оферта, проверка требований к платёжной деятельности. Допущение.'); r += 1
inp(r, 'acct', 'Бухгалтерия и юрист, в месяц', 40000, TENGE, '₸/мес', 'Допущение: аутсорс для ИП.'); r += 1
r += 1
A.cell(r, 2, 'По годам').font = H2
for i, y in enumerate(['Год 1', 'Год 2', 'Год 3']):
    A.cell(r, 3 + i, y).font = BOLD
r += 1

def inp3(row, key, label, vals, fmt, unit, note, key_param=False):
    A.cell(row, 2, label)
    for i, v in enumerate(vals):
        c = A.cell(row, 3 + i, v); c.font = BLUE; c.number_format = fmt
        if key_param: c.fill = YELLOW
        names[f'{key}{i+1}'] = f"Допущения!${L(3+i)}${row}"
    A.cell(row, 6, unit).font = MUTED
    A.cell(row, 7, note).font = MUTED

inp3(r, 'host', 'Хостинг и сервисы, в месяц', [60000, 120000, 200000], TENGE, '₸/мес', 'Допущение: VPS + управляемые PostgreSQL/Redis, рост с нагрузкой.'); r += 1
inp3(r, 'salary', 'Выплата основателю, в месяц', [0, 300000, 500000], TENGE, '₸/мес', 'Бизнес-решение: команда — только основатель. Суммы выплат — допущение, уточнить.', True); r += 1
r += 1

A.cell(r, 2, 'Сценарии').font = H2
SC = ['Пессимистичный', 'Базовый', 'Оптимистичный']
for i, s in enumerate(SC):
    A.cell(r, 3 + i, s).font = BOLD
r += 1
inp3(r, 'price', 'Средняя цена билета', [5000, 6000, 7000], TENGE, '₸', 'МСБ-сегмент дешевле крупных концертов. Ориентир: средний чек Ticketon 2021 — $11,9 ≈ 5 700 ₸ (spec.md).', True); r += 1
inp3(r, 'acq', 'Эквайринг (от суммы оплаты)', [0.025, 0.02, 0.015], PCT, 'от оплаты', 'Kaspi Pay ~1–1,5%, Halyk ePay 2,5–3% для карт других банков (finkaz.kz, halykbank.kz, 2026).'); r += 1
inp3(r, 'new1', 'Новых организаторов в месяц, год 1', [1.5, 3, 5], NUM1, 'шт./мес', 'Прямые продажи основателя + бесплатное первое событие (spec.md: первые 20–50 организаторов).', True); r += 1
inp3(r, 'new2', 'Новых организаторов в месяц, год 2', [3, 5, 8], NUM1, 'шт./мес', 'Рекомендации, второй город. Допущение.', True); r += 1
inp3(r, 'new3', 'Новых организаторов в месяц, год 3', [4, 7, 11], NUM1, 'шт./мес', 'Другие города Казахстана. Допущение.', True); r += 1
inp3(r, 'churn', 'Отток организаторов в месяц', [0.06, 0.045, 0.035], PCT, 'от активных', 'Допущение: для SaaS МСБ обычно 3–7% в месяц.'); r += 1
inp3(r, 'epo', 'Событий на организатора в месяц', [1.0, 1.2, 1.5], NUM1, 'шт.', 'Стендап-клуб — еженедельно, конференция — раз в квартал; среднее. Допущение.'); r += 1
inp3(r, 'tpe', 'Продано билетов на событие', [90, 110, 150], NUM1, 'шт.', 'Малые и средние залы. Допущение.'); r += 1
inp3(r, 'cac', 'Стоимость привлечения организатора (CAC)', [200000, 150000, 100000], TENGE, '₸', 'Реклама, встречи, время на продажу, бесплатное первое событие. Допущение: для B2B-продаж МСБ дороже, чем кажется.'); r += 1
# Сценарные ключи: price1 = пессимистичный, price2 = базовый, price3 = оптимистичный.
r += 1
A.cell(r, 2, 'Сезонность (1,0 — обычный месяц)').font = H2
r += 1
season = [0.8, 0.9, 1.0, 1.0, 1.0, 0.8, 0.6, 0.6, 1.0, 1.1, 1.2, 1.3]
months_ru = ['янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек']
season_row = r + 1
for i, (m, v) in enumerate(zip(months_ru, season)):
    A.cell(r, 3 + i, m).font = BOLD
    c = A.cell(r + 1, 3 + i, v); c.font = BLUE; c.number_format = '0.00'
A.cell(r + 1, 2, 'Коэффициент')
A.cell(r + 2, 2, 'Летом залы пустеют, к Новому году — пик. Допущение.').font = MUTED
season_ref = f"Допущения!$C${season_row}:$N${season_row}"
r += 4
A.cell(r, 2, 'Крупный сегмент: концерты и шоу от 2 000 мест').font = H2
for i, s in enumerate(SC):
    A.cell(r, 3 + i, s).font = BOLD
A.cell(r, 7, 'Бизнес-решение владельца 2026-10-09: оба сегмента в финмодели. Все числа ниже — допущения без интервью: '
             'по решению владельца 2026-10-09 интервью до защиты не проводятся. Источники — лист «Риски и источники».').font = MUTED
r += 1
inp3(r, 'lev1_', 'Крупных событий в год, год 1', [0, 0, 1], NUM1, 'шт./год',
     'Допущение. Запуск в июне 2027, первые контакты с промоутерами — после запуска. Промоутеры связаны с банковскими '
     'площадками (Ticketon у Freedom, kino.kz у Halyk) и эксклюзивами. Оптимистичный — один пилот.', True); r += 1
inp3(r, 'lev2_', 'Крупных событий в год, год 2', [0, 1, 2], NUM1, 'шт./год',
     'Допущение. Реалистично — пилот с местным артистом или квота мест на арене, а не тур мировой звезды.', True); r += 1
inp3(r, 'lev3_', 'Крупных событий в год, год 3', [1, 2, 4], NUM1, 'шт./год',
     'Допущение. Событий арен и стадионов около 60 в год (лист «Рынок»): база 2 ≈ 3%, оптимистичный 4 ≈ 7%. '
     'Для одного основателя это потолок.', True); r += 1
inp3(r, 'ltpe', 'Продано билетов на крупное событие', [4000, 7000, 12000], TENGE, 'шт.',
     'Допущение. Almaty Arena: 12 концертов и 78 950 зрителей за I полугодие 2024 (акимат Алматы) ≈ 6 600 на концерт; '
     'Би-2 — около 8 000 (Informburo, 2022). Пессимистичный — малая арена или квота мест, оптимистичный — полная арена. Целый стадион не закладываем.', True); r += 1
inp3(r, 'lprice', 'Средняя цена билета на крупном событии', [12000, 15000, 20000], TENGE, '₸',
     'Допущение. Средний чек на концерты за январь–июль 2024 — 28 232 ₸ (Ticketon по данным bes.media); из источника неясно, на заказ это '
     'или на билет. Если на заказ из 2 билетов — около 14 000 ₸ за билет; если на билет — базовые 15 000 ₸ консервативны. '
     'Макс Корж на стадионе — от 23 000 ₸ (концерт 2025 отменён площадкой), Skillet — 20 000–50 000 ₸. Цены мировых звёзд (40 000–350 000 ₸) не закладываем.', True); r += 1
inp3(r, 'lcost', 'Продажа и сопровождение одного крупного события', [2000000, 1000000, 600000], TENGE, '₸',
     'Допущение. База: переговоры и юрист 200–300 тыс., нагрузочный прогон и запас облака на пик ~100 тыс., дежурство на старте '
     'продаж и в день события 150–250 тыс., поддержка покупателей, уступки промоутеру 300–400 тыс. ₸. Пессимистичный — первое событие: дольше переговоры, больше уступки и запас облака; '
     'оптимистичный — повторный промоутер, отлаженный процесс. На билет — 500 / 143 / 50 ₸.', True); r += 1
A.cell(r, 2, 'Сбор, комиссия, налог, билетов в заказе, SMS, инфраструктура, доли бесплатных и возвратов — общие с МСБ; '
             'эквайринг — по сценарию. События распределены по году равномерно, без сезонности: доля события попадает в каждый месяц. '
             'Поэтому месяцы первой прибыли и окупаемости на листе «Итоги» в базовом и оптимистичном сценариях раньше реальных: '
             'пилот 1-го года даёт вклад с 1-го месяца, хотя пройдёт через несколько месяцев после запуска. Суммы за год и за 3 года верны.').font = MUTED
for col in range(8, 15):
    A.column_dimensions[L(col)].width = 7
A.column_dimensions['G'].width = 90
A.freeze_panes = 'C5'

# ---------------------------------------------------------------- План
P = wb.create_sheet('План')
P['A1'] = 'План на 36 месяцев по трём сценариям'; P['A1'].font = TITLE
P['A2'] = ('Все ячейки — формулы от листа «Допущения». Суммы в тенге. В каждом сценарии сначала МСБ, затем крупный сегмент; '
           'операционная прибыль — оба сегмента.'); P['A2'].font = MUTED
P.column_dimensions['A'].width = 44
P.column_dimensions['B'].width = 9
M0 = 3; N = 36
cols = [L(M0 + i) for i in range(N)]
for c in cols:
    P.column_dimensions[c].width = 12
P.cell(4, 1, 'Месяц').font = BOLD
P.cell(5, 1, 'Дата').font = BOLD
P.cell(6, 1, 'Год модели').font = BOLD
P.cell(7, 1, 'Сезонность').font = BOLD
for i, c in enumerate(cols):
    P[f'{c}4'] = i + 1
    P[f'{c}5'] = f'=EDATE({names["start"]},{i})'; P[f'{c}5'].number_format = 'mmm yy'
    P[f'{c}6'] = f'=INT(({c}4-1)/12)+1'
    P[f'{c}7'] = f'=INDEX({season_ref},MONTH({c}5))'; P[f'{c}7'].number_format = '0.00'
    for rr in (4, 5, 6, 7):
        P[f'{c}{rr}'].font = BOLD if rr == 4 else BLACK
        P[f'{c}{rr}'].fill = HEAD
for rr in (4, 5, 6, 7):
    P.cell(rr, 1).fill = HEAD; P.cell(rr, 2).fill = HEAD

LARGE = ['l_events', 'l_tickets', 'l_paid', 'l_gmv', 'l_rev', 'l_acq', 'l_var', 'l_tax', 'l_margin', 'l_cost']
ROWS = ['new', 'churned', 'active', 'events', 'tickets', 'paid', 'gmv', 'rev', 'acq', 'var', 'tax', 'margin',
        *LARGE, 'mkt', 'host', 'acct', 'salary', 'fixed', 'profit', 'cash', 'pflag', 'cflag']
LABELS = {
    'new': ('Новые организаторы МСБ', 'шт.', NUM1),
    'churned': ('Ушедшие организаторы МСБ', 'шт.', NUM1),
    'active': ('Активные организаторы МСБ на конец месяца', 'шт.', NUM1),
    'events': ('Событий с продажей билетов (МСБ)', 'шт.', NUM1),
    'tickets': ('Продано билетов (МСБ)', 'шт.', TENGE),
    'paid': ('Платных билетов (МСБ)', 'шт.', TENGE),
    'gmv': ('Оборот билетов организаторов (GMV, МСБ)', '₸', TENGE),
    'rev': ('Выручка: сервисный сбор и комиссия (МСБ)', '₸', TENGE),
    'acq': ('Эквайринг (МСБ)', '₸', TENGE),
    'var': ('SMS и инфраструктура на билеты (МСБ)', '₸', TENGE),
    'tax': ('Налог с выручки (МСБ)', '₸', TENGE),
    'margin': ('Маржинальная прибыль (МСБ)', '₸', TENGE),
    'l_events': ('Крупный сегмент: событий', 'шт.', NUM1),
    'l_tickets': ('Крупный сегмент: продано билетов', 'шт.', TENGE),
    'l_paid': ('Крупный сегмент: платных билетов', 'шт.', TENGE),
    'l_gmv': ('Крупный сегмент: оборот билетов (GMV)', '₸', TENGE),
    'l_rev': ('Крупный сегмент: выручка, сбор и комиссия', '₸', TENGE),
    'l_acq': ('Крупный сегмент: эквайринг', '₸', TENGE),
    'l_var': ('Крупный сегмент: SMS и инфраструктура', '₸', TENGE),
    'l_tax': ('Крупный сегмент: налог с выручки', '₸', TENGE),
    'l_margin': ('Крупный сегмент: маржинальная прибыль', '₸', TENGE),
    'l_cost': ('Крупный сегмент: продажа и сопровождение событий', '₸', TENGE),
    'mkt': ('Привлечение организаторов МСБ (CAC × новые)', '₸', TENGE),
    'host': ('Хостинг и сервисы', '₸', TENGE),
    'acct': ('Бухгалтерия, юрист, запуск', '₸', TENGE),
    'salary': ('Выплата основателю', '₸', TENGE),
    'fixed': ('Постоянные расходы и маркетинг', '₸', TENGE),
    'profit': ('Операционная прибыль', '₸', TENGE),
    'cash': ('Накопленный денежный поток', '₸', TENGE),
    'pflag': ('служебная: месяц с прибылью', '', '0'),
    'cflag': ('служебная: месяц с накопленным потоком ≥ 0', '', '0'),
}
block = {}
row = 9
for si, sname in enumerate(SC, start=1):
    P.cell(row, 1, f'Сценарий: {sname}').font = H2
    rows = {k: row + 1 + j for j, k in enumerate(ROWS)}
    block[si] = rows
    S = lambda k: names[f'{k}{si}']
    for k in ROWS:
        rr = rows[k]
        lab, unit, fmt = LABELS[k]
        P.cell(rr, 1, lab).font = MUTED if k in ('pflag', 'cflag') else (BOLD if k in ('margin', 'l_margin', 'profit', 'cash') else BLACK)
        P.cell(rr, 2, unit).font = MUTED
        for i, c in enumerate(cols):
            prev = cols[i - 1] if i else None
            y = f'{c}$6'
            f = {
                'new': f'=CHOOSE({y},{S("new1")},{S("new2")},{S("new3")})',
                'churned': f'={prev}{rows["active"]}*{S("churn")}' if prev else '=0',
                'active': (f'={prev}{rows["active"]}+{c}{rows["new"]}-{c}{rows["churned"]}' if prev else f'={c}{rows["new"]}'),
                'events': f'={c}{rows["active"]}*{S("epo")}*{c}$7',
                'tickets': f'={c}{rows["events"]}*{S("tpe")}',
                'paid': f'={c}{rows["tickets"]}*(1-{names["free"]})',
                'gmv': f'={c}{rows["paid"]}*{S("price")}',
                'rev': f'={c}{rows["gmv"]}*({names["fee"]}+{names["orgfee"]})*(1-{names["refund"]})',
                'acq': f'={c}{rows["gmv"]}*(1+{names["fee"]})*{S("acq")}',
                'var': f'={c}{rows["tickets"]}/{names["tpo"]}*{names["sms"]}+{c}{rows["tickets"]}*{names["infra"]}',
                'tax': f'={c}{rows["rev"]}*{names["tax"]}',
                'margin': f'={c}{rows["rev"]}-{c}{rows["acq"]}-{c}{rows["var"]}-{c}{rows["tax"]}',
                'l_events': f'=CHOOSE({y},{S("lev1_")},{S("lev2_")},{S("lev3_")})/12',
                'l_tickets': f'={c}{rows["l_events"]}*{S("ltpe")}',
                'l_paid': f'={c}{rows["l_tickets"]}*(1-{names["free"]})',
                'l_gmv': f'={c}{rows["l_paid"]}*{S("lprice")}',
                'l_rev': f'={c}{rows["l_gmv"]}*({names["fee"]}+{names["orgfee"]})*(1-{names["refund"]})',
                'l_acq': f'={c}{rows["l_gmv"]}*(1+{names["fee"]})*{S("acq")}',
                'l_var': f'={c}{rows["l_tickets"]}/{names["tpo"]}*{names["sms"]}+{c}{rows["l_tickets"]}*{names["infra"]}',
                'l_tax': f'={c}{rows["l_rev"]}*{names["tax"]}',
                'l_margin': f'={c}{rows["l_rev"]}-{c}{rows["l_acq"]}-{c}{rows["l_var"]}-{c}{rows["l_tax"]}',
                'l_cost': f'={c}{rows["l_events"]}*{S("lcost")}',
                'mkt': f'={c}{rows["new"]}*{S("cac")}',
                'host': f'=CHOOSE({y},{names["host1"]},{names["host2"]},{names["host3"]})',
                'acct': f'={names["acct"]}+IF({c}$4=1,{names["legal1"]},0)',
                'salary': f'=CHOOSE({y},{names["salary1"]},{names["salary2"]},{names["salary3"]})',
                'fixed': f'=SUM({c}{rows["mkt"]}:{c}{rows["salary"]})',
                'profit': f'={c}{rows["margin"]}+{c}{rows["l_margin"]}-{c}{rows["fixed"]}-{c}{rows["l_cost"]}',
                'cash': (f'={prev}{rows["cash"]}+{c}{rows["profit"]}' if prev else f'={c}{rows["profit"]}'),
                'pflag': f'=IF({c}{rows["profit"]}>0,{c}$4,"")',
                'cflag': f'=IF({c}{rows["cash"]}>=0,{c}$4,"")',
            }[k]
            cell = P[f'{c}{rr}']
            cell.value = f
            cell.number_format = fmt
            cell.font = MUTED if k in ('pflag', 'cflag') else (BOLD if k in ('margin', 'l_margin', 'profit', 'cash') else BLACK)
        if k in ('margin', 'l_margin', 'profit', 'cash'):
            for i in range(N + 2):
                P.cell(rr, 1 + i).border = THIN
    row = rows['cflag'] + 3
P.freeze_panes = 'C8'

# ---------------------------------------------------------------- Итоги
I = wb.create_sheet('Итоги', 0)
I['A1'] = 'Финансовая модель «Партер»: итоги за три года'; I['A1'].font = TITLE
I['A2'] = ('Два сегмента: организаторы МСБ и крупные концерты и шоу от 2 000 мест (оба в финмодели — бизнес-решение 2026-10-09). '
           'Модель — сервисный сбор с покупателя (бизнес-решение 2026-10-02), организатор не платит; команда — только основатель; '
           'Казахстан, тенге. Строки 6–12 и ключевые выводы, кроме LTV/CAC, — оба сегмента вместе, крупный сегмент отдельно — ниже. '
           'Крупные события размазаны по месяцам равномерно, поэтому месяцы первой прибыли и окупаемости (строки 15–16) '
           'раньше реальных: в оптимистичном сценарии без крупного сегмента — 4-й и 7-й, в базовом первая прибыль — 31-й месяц. '
           'Все значения — формулы от листа «Допущения».')
I['A2'].font = MUTED
I.column_dimensions['A'].width = 62
for col in range(2, 11):
    I.column_dimensions[L(col)].width = 15

def yrange(si, key, y):
    rr = block[si][key]
    a, b = cols[(y - 1) * 12], cols[y * 12 - 1]
    return f"План!${a}${rr}:${b}${rr}"

def rowref(si, key):
    rr = block[si][key]
    return f"План!${cols[0]}${rr}:${cols[-1]}${rr}"

I.cell(4, 1, 'Показатель').font = BOLD
hdr = []
for si, s in enumerate(SC, start=1):
    for y in (1, 2, 3):
        hdr.append((si, y))
# Компактная таблица: строки — показатели, столбцы — сценарий × год
I.cell(3, 2, 'Пессимистичный').font = BOLD
I.cell(3, 5, 'Базовый').font = BOLD
I.cell(3, 8, 'Оптимистичный').font = BOLD
for j, (si, y) in enumerate(hdr):
    c = I.cell(4, 2 + j, f'Год {y}'); c.font = BOLD; c.fill = HEAD
I.cell(4, 1).fill = HEAD
# Ключи через «+» складываются: МСБ + крупный сегмент.
metrics = [
    ('Активных организаторов МСБ на конец года', 'active', 'end', NUM1),
    ('Продано билетов', 'tickets+l_tickets', 'sum', TENGE),
    ('Оборот билетов (GMV), ₸', 'gmv+l_gmv', 'sum', TENGE),
    ('Выручка, ₸', 'rev+l_rev', 'sum', TENGE),
    ('Маржинальная прибыль, ₸', 'margin+l_margin', 'sum', TENGE),
    ('Постоянные расходы, маркетинг и продажа крупных событий, ₸', 'fixed+l_cost', 'sum', TENGE),
    ('Операционная прибыль, ₸', 'profit', 'sum', TENGE),
    ('Накопленный поток на конец года, ₸', 'cash', 'end', TENGE),
]

def metric_formula(si, y, keys, agg):
    if agg == 'sum':
        return '=' + '+'.join(f'SUM({yrange(si, k, y)})' for k in keys.split('+'))
    return f"=План!{cols[y * 12 - 1]}{block[si][keys]}"

rr = 5
for lab, key, agg, fmt in metrics:
    I.cell(rr, 1, lab).font = BOLD if key in ('profit', 'cash') else BLACK
    for j, (si, y) in enumerate(hdr):
        c = I.cell(rr, 2 + j, metric_formula(si, y, key, agg)); c.number_format = fmt; c.font = GREEN
    rr += 1
rr += 1
I.cell(rr, 1, 'Ключевые выводы').font = H2
for j, s in enumerate(SC):
    I.cell(rr, 2 + j * 3, s).font = BOLD
rr += 1
key_rows = [
    ('Первый месяц с операционной прибылью', lambda si: f'=IF(COUNT({rowref(si, "pflag")})=0,"не достигнут",MIN({rowref(si, "pflag")}))', '0'),
    ('Месяц, с которого накопленный поток ≥ 0 (окупаемость)', lambda si: f'=IF(COUNT({rowref(si, "cflag")})=0,"за 3 года нет",MIN({rowref(si, "cflag")}))', '0'),
    ('Нужно вложить до окупаемости (минимум накопленного потока), ₸', lambda si: f'=-MIN(0,MIN({rowref(si, "cash")}))', TENGE),
    ('Выручка за 3 года, ₸', lambda si: f'=SUM({rowref(si, "rev")})+SUM({rowref(si, "l_rev")})', TENGE),
    ('Операционная прибыль за 3 года, ₸', lambda si: f'=SUM({rowref(si, "profit")})', TENGE),
    ('LTV / CAC (организатор МСБ)', lambda si: f"='Юнит-экономика'!{L(2 + si)}21", '0.0"x"'),
]
key_row_of = {}
for lab, fn, fmt in key_rows:
    I.cell(rr, 1, lab)
    key_row_of[lab] = rr
    for si in (1, 2, 3):
        c = I.cell(rr, 2 + (si - 1) * 3, fn(si)); c.number_format = fmt; c.font = GREEN; c.alignment = Alignment(horizontal='right')
    rr += 1
profit_row = key_row_of['Операционная прибыль за 3 года, ₸']

# В том числе крупный сегмент: те же 9 столбцов сценарий × год.
rr += 1
I.cell(rr, 1, 'В том числе крупный сегмент').font = H2
for j, s_ in enumerate(SC):
    I.cell(rr, 2 + j * 3, s_).font = BOLD
rr += 1
I.cell(rr, 1, 'Показатель').font = BOLD; I.cell(rr, 1).fill = HEAD
for j, (si, y) in enumerate(hdr):
    c = I.cell(rr, 2 + j, f'Год {y}'); c.font = BOLD; c.fill = HEAD
rr += 1
large_metrics = [
    ('Крупных событий', 'l_events', NUM1),
    ('Продано билетов', 'l_tickets', TENGE),
    ('Выручка, ₸', 'l_rev', TENGE),
    ('Маржинальная прибыль, ₸', 'l_margin', TENGE),
    ('Продажа и сопровождение событий, ₸', 'l_cost', TENGE),
    ('Вклад в операционную прибыль, ₸', None, TENGE),
]
for lab, key, fmt in large_metrics:
    I.cell(rr, 1, lab).font = BOLD if key is None else BLACK
    for j, (si, y) in enumerate(hdr):
        if key is None:
            f = f'=SUM({yrange(si, "l_margin", y)})-SUM({yrange(si, "l_cost", y)})'
        else:
            f = f'=SUM({yrange(si, key, y)})'
        c = I.cell(rr, 2 + j, f); c.number_format = fmt; c.font = BOLD if key is None else GREEN
    contrib_row = rr
    rr += 1
rr += 1
I.cell(rr, 1, 'Операционная прибыль за 3 года без крупного сегмента, ₸')
for si in (1, 2, 3):
    a = 2 + (si - 1) * 3
    f = f'={L(a)}{profit_row}-SUM({L(a)}{contrib_row}:{L(a + 2)}{contrib_row})'
    c = I.cell(rr, a, f); c.number_format = TENGE; c.alignment = Alignment(horizontal='right')
rr += 1
chart_anchor_row = rr + 2

# График накопленного потока
ch = LineChart()
ch.title = 'Накопленный денежный поток, ₸'
ch.y_axis.title = '₸'
ch.x_axis.title = 'Месяц'
ch.height, ch.width = 9, 24
for si, color in zip((1, 2, 3), ('E3494B', '2A78D6', '1BAF7A')):
    rr_ = block[si]['cash']
    data = Reference(P, min_col=M0, max_col=M0 + N - 1, min_row=rr_, max_row=rr_)
    ch.add_data(data, from_rows=True, titles_from_data=False)
    s = ch.series[-1]
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 22000
    s.smooth = False
from openpyxl.chart.series import SeriesLabel
for s, name in zip(ch.series, SC):
    s.tx = SeriesLabel(v=name)
ch.set_categories(Reference(P, min_col=M0, max_col=M0 + N - 1, min_row=4, max_row=4))
ch.y_axis.number_format = '#,##0'
ch.y_axis.majorGridlines.spPr = None
I.add_chart(ch, f'A{chart_anchor_row}')

# ---------------------------------------------------------------- Юнит-экономика
U = wb.create_sheet('Юнит-экономика', 1)
U['A1'] = 'Юнит-экономика: один билет, один организатор МСБ и одно крупное событие'; U['A1'].font = TITLE
U['A2'] = 'Цены и ставки — из листа «Допущения»; маржа на билет учитывает бесплатные билеты и возвраты.'; U['A2'].font = MUTED
U.column_dimensions['A'].width = 4; U.column_dimensions['B'].width = 58
for col in 'CDE': U.column_dimensions[col].width = 18
U.column_dimensions['F'].width = 60
for i, s in enumerate(SC):
    c = U.cell(4, 3 + i, s); c.font = BOLD; c.fill = HEAD
U.cell(4, 2, 'На один проданный билет (МСБ)').font = H2; U.cell(4, 2).fill = HEAD
u_rows = [
    (5, 'Цена билета, ₸', lambda si: f'={names[f"price{si}"]}', TENGE, ''),
    (6, 'Сервисный сбор и комиссия с платного билета, ₸', lambda si: f'=C5*0+{L(2+si)}5*({names["fee"]}+{names["orgfee"]})', TENGE, ''),
    (7, 'Средняя выручка на проданный билет, ₸', lambda si: f'={L(2+si)}6*(1-{names["free"]})*(1-{names["refund"]})', TENGE, 'с учётом бесплатных билетов и возвратов'),
    (8, 'Эквайринг на проданный билет, ₸', lambda si: f'={L(2+si)}5*(1+{names["fee"]})*{names[f"acq{si}"]}*(1-{names["free"]})', TENGE, 'с суммы оплаты, включая сбор; при возврате не возвращается'),
    (9, 'SMS и инфраструктура на билет, ₸', lambda si: f'={names["sms"]}/{names["tpo"]}+{names["infra"]}', TENGE, ''),
    (10, 'Налог, ₸', lambda si: f'={L(2+si)}7*{names["tax"]}', TENGE, ''),
    (11, 'Маржа на проданный билет, ₸', lambda si: f'={L(2+si)}7-{L(2+si)}8-{L(2+si)}9-{L(2+si)}10', TENGE, ''),
    (12, 'Маржа, % от цены билета', lambda si: f'=IF({L(2+si)}5=0,0,{L(2+si)}11/{L(2+si)}5)', PCT, 'в первой версии spec.md было ~2% при сборе 5% и эквайринге 2,5%'),
]
for rr_, lab, fn, fmt, note in u_rows:
    U.cell(rr_, 2, lab).font = BOLD if rr_ in (11, 12) else BLACK
    for si in (1, 2, 3):
        c = U.cell(rr_, 2 + si, fn(si)); c.number_format = fmt; c.font = GREEN if rr_ == 5 else (BOLD if rr_ in (11, 12) else BLACK)
    U.cell(rr_, 6, note).font = MUTED
# фикс строки 6: убрать артефакт C5*0
for si in (1, 2, 3):
    U.cell(6, 2 + si).value = f'={L(2+si)}5*({names["fee"]}+{names["orgfee"]})'
U.cell(14, 2, 'На одного организатора МСБ').font = H2; U.cell(14, 2).fill = HEAD
for i in range(3): U.cell(14, 3 + i).fill = HEAD
o_rows = [
    (15, 'Билетов в месяц (в среднем по году, с сезонностью)', lambda si: f'={names[f"epo{si}"]}*{names[f"tpe{si}"]}*AVERAGE({season_ref})', NUM1, ''),
    (16, 'Маржа с организатора в месяц, ₸', lambda si: f'={L(2+si)}15*{L(2+si)}11', TENGE, ''),
    (17, 'Отток в месяц', lambda si: f'={names[f"churn{si}"]}', PCT, ''),
    (18, 'Средний срок жизни организатора, мес', lambda si: f'=IF({L(2+si)}17=0,0,1/{L(2+si)}17)', NUM1, '1 / отток'),
    (19, 'LTV (маржа за срок жизни), ₸', lambda si: f'={L(2+si)}16*{L(2+si)}18', TENGE, ''),
    (20, 'CAC, ₸', lambda si: f'={names[f"cac{si}"]}', TENGE, ''),
    (21, 'LTV / CAC', lambda si: f'=IF({L(2+si)}20=0,0,{L(2+si)}19/{L(2+si)}20)', '0.0"x"', 'больше 3x — привлечение окупается с запасом'),
    (22, 'Окупаемость привлечения, мес', lambda si: f'=IF({L(2+si)}16=0,0,{L(2+si)}20/{L(2+si)}16)', NUM1, 'CAC / маржа в месяц'),
    (23, 'Организаторов для безубыточности во 2-м году', lambda si: f'=IF({L(2+si)}16=0,0,({names["host2"]}+{names["acct"]}+{names["salary2"]})/{L(2+si)}16)', NUM1, 'постоянные расходы 2-го года без маркетинга / маржа с организатора'),
]
for rr_, lab, fn, fmt, note in o_rows:
    U.cell(rr_, 2, lab).font = BOLD if rr_ in (19, 21, 23) else BLACK
    for si in (1, 2, 3):
        c = U.cell(rr_, 2 + si, fn(si)); c.number_format = fmt
        c.font = GREEN if rr_ in (17, 20) else (BOLD if rr_ in (19, 21, 23) else BLACK)
    U.cell(rr_, 6, note).font = MUTED

# На одно крупное событие: те же формулы, что в строках 5–11, плюс затраты на событие.
U.cell(25, 2, 'На одно крупное событие').font = H2; U.cell(25, 2).fill = HEAD
for i, s_ in enumerate(SC):
    c = U.cell(25, 3 + i, s_); c.font = BOLD; c.fill = HEAD
X = lambda si: L(2 + si)
l_rows = [
    (26, 'Цена билета, ₸', lambda si: f'={names[f"lprice{si}"]}', TENGE, ''),
    (27, 'Сервисный сбор и комиссия с платного билета, ₸', lambda si: f'={X(si)}26*({names["fee"]}+{names["orgfee"]})', TENGE, ''),
    (28, 'Средняя выручка на проданный билет, ₸', lambda si: f'={X(si)}27*(1-{names["free"]})*(1-{names["refund"]})', TENGE, 'с учётом бесплатных билетов и возвратов'),
    (29, 'Эквайринг на проданный билет, ₸', lambda si: f'={X(si)}26*(1+{names["fee"]})*{names[f"acq{si}"]}*(1-{names["free"]})', TENGE, 'ставка эквайринга сценария, как у МСБ'),
    (30, 'SMS и инфраструктура на билет, ₸', lambda si: f'={names["sms"]}/{names["tpo"]}+{names["infra"]}', TENGE, ''),
    (31, 'Налог, ₸', lambda si: f'={X(si)}28*{names["tax"]}', TENGE, ''),
    (32, 'Маржа на проданный билет, ₸', lambda si: f'={X(si)}28-{X(si)}29-{X(si)}30-{X(si)}31', TENGE, ''),
    (33, 'Продано билетов на событие', lambda si: f'={names[f"ltpe{si}"]}', TENGE, ''),
    (34, 'Маржа с события, ₸', lambda si: f'={X(si)}32*{X(si)}33', TENGE, ''),
    (35, 'Продажа и сопровождение события, ₸', lambda si: f'={names[f"lcost{si}"]}', TENGE, 'переговоры, юрист, запас облака, дежурство, поддержка, уступки промоутеру'),
    (36, 'Вклад события в операционную прибыль, ₸', lambda si: f'={X(si)}34-{X(si)}35', TENGE, ''),
    (37, 'Точка безубыточности, билетов на событие', lambda si: f'=IF({X(si)}32<=0,"нет",{X(si)}35/{X(si)}32)', TENGE, 'затраты на событие / маржа на билет'),
]
for rr_, lab, fn, fmt, note in l_rows:
    bold = rr_ in (32, 36, 37)
    U.cell(rr_, 2, lab).font = BOLD if bold else BLACK
    for si in (1, 2, 3):
        c = U.cell(rr_, 2 + si, fn(si)); c.number_format = fmt
        c.font = GREEN if rr_ in (26, 33, 35) else (BOLD if bold else BLACK)
        if rr_ == 37: c.alignment = Alignment(horizontal='right')
    U.cell(rr_, 6, note).font = MUTED

# ---------------------------------------------------------------- Сравнение моделей
C = wb.create_sheet('Сравнение моделей')
C['A1'] = 'Почему процент с билета, а не подписка'; C['A1'].font = TITLE
C['A2'] = ('Только организаторы МСБ: объёмы МСБ базового сценария 3-го года из листа «План»; меняется только способ брать деньги. '
           'Подписку и гибрид предлагала исходная spec.md; выбран сервисный сбор (бизнес-решение 2026-10-02).')
C['A2'].font = MUTED
C.column_dimensions['A'].width = 4; C.column_dimensions['B'].width = 56
for col in 'CDEF': C.column_dimensions[col].width = 20
C.column_dimensions['G'].width = 50
C.cell(4, 2, 'Вводные для сравнения').font = H2
C.cell(5, 2, 'Подписка организатора, ₸ в месяц'); C['C5'] = 25000; C['C5'].font = BLUE; C['C5'].number_format = TENGE; C['C5'].fill = YELLOW
C['G5'] = 'Альтернатива из spec.md; цена — допущение'; C['G5'].font = MUTED
C.cell(6, 2, 'Гибрид: подписка, ₸ в месяц'); C['C6'] = 15000; C['C6'].font = BLUE; C['C6'].number_format = TENGE
C.cell(7, 2, 'Гибрид: комиссия с билета'); C['C7'] = 0.01; C['C7'].font = BLUE; C['C7'].number_format = PCT
C['G7'] = 'spec.md: подписка + 1% с билета'; C['G7'].font = MUTED
C.cell(8, 2, 'Доля организаторов, готовых платить подписку'); C['C8'] = 0.5; C['C8'].font = BLUE; C['C8'].number_format = PCT; C['C8'].fill = YELLOW
C['G8'] = 'Допущение: МСБ с редкими событиями не платит каждый месяц; процент без абонплаты берут все'; C['G8'].font = MUTED

C.cell(10, 2, 'Объёмы организаторов МСБ, базовый сценарий, 3-й год').font = H2
orgm = f'SUM({yrange(2, "active", 3)})'
C.cell(11, 2, 'Организаторо-месяцев'); C['C11'] = f'={orgm}'; C['C11'].number_format = TENGE; C['C11'].font = GREEN
C.cell(12, 2, 'Платных билетов'); C['C12'] = f'=SUM({yrange(2, "paid", 3)})'; C['C12'].number_format = TENGE; C['C12'].font = GREEN
C.cell(13, 2, 'Оборот билетов (GMV), ₸'); C['C13'] = f'=SUM({yrange(2, "gmv", 3)})'; C['C13'].number_format = TENGE; C['C13'].font = GREEN

for i, h in enumerate(['Выручка, ₸', 'Выручка на организатора в месяц, ₸', 'Кто платит']):
    c = C.cell(15, 3 + i, h); c.font = BOLD; c.fill = HEAD
C.cell(15, 2, 'Модель').font = BOLD; C.cell(15, 2).fill = HEAD
ref = (1 - 0.0)
models = [
    ('Сервисный сбор 5% (выбрано)', f'=C13*{names["fee"]}*(1-{names["refund"]})', 'покупатель'),
    ('Подписка', '=C11*C8*C5', 'организатор'),
    ('Гибрид: подписка + 1%', f'=C11*C8*C6+C13*C7*(1-{names["refund"]})', 'организатор'),
]
for j, (name, f, who) in enumerate(models):
    rr_ = 16 + j
    C.cell(rr_, 2, name).font = BOLD if j == 0 else BLACK
    C.cell(rr_, 3, f).number_format = TENGE
    C.cell(rr_, 4, f'=IF($C$11=0,0,C{rr_}/$C$11)').number_format = TENGE
    C.cell(rr_, 5, who)
C.cell(20, 2, 'Точка равенства').font = H2
C.cell(21, 2, 'Платных билетов в месяц у организатора, при которых сбор = подписке')
C['C21'] = f'=IF({names["price2"]}*{names["fee"]}*(1-{names["refund"]})=0,0,C5/({names["price2"]}*{names["fee"]}*(1-{names["refund"]})))'
C['C21'].number_format = NUM1
C['G21'] = 'меньше — организатору выгоднее процент, больше — подписка; средний организатор базового сценария (строка 22) выше точки равенства'; C['G21'].font = MUTED
C.cell(22, 2, 'Платных билетов в месяц у организатора в базовом сценарии')
C['C22'] = f"='Юнит-экономика'!D15*(1-{names['free']})"; C['C22'].number_format = NUM1; C['C22'].font = GREEN
C.cell(24, 2, 'Вывод').font = H2
C['B25'] = ('Сервисный сбор платит покупатель, поэтому организатору вход бесплатный — это главный аргумент для МСБ против '
            'крупных площадок, где обычно платят обе стороны (схема «10/10» у Ticketscloud: сбор с покупателя и комиссия с организатора). Подписка выгоднее организаторам '
            'с объёмом выше точки равенства, а средний организатор базового сценария — около 112 платных билетов в месяц — как раз выше неё. '
            'Поэтому сбор приносит больше подписки даже с каждого платящего организатора (примерно в 1,3 раза), а то, что подписку платит '
            'только половина организаторов, удваивает разрыв.')
C['B25'].alignment = Alignment(wrap_text=True, vertical='top')
C.merge_cells('B25:G28')

# ---------------------------------------------------------------- Рынок
R = wb.create_sheet('Рынок')
R['A1'] = 'Рынок снизу вверх: МСБ-организаторы и крупный сегмент'; R['A1'].font = TITLE
R['A2'] = ('Все численности — допущения для проверки порядка величин. Интервью до защиты не проводятся (решение владельца '
           '2026-10-09); численности проверяются открытыми данными и данными первых месяцев после запуска.'); R['A2'].font = MUTED
R.column_dimensions['A'].width = 4; R.column_dimensions['B'].width = 46
for col in 'CDEFGH': R.column_dimensions[col].width = 16
heads = ['Организаторов', 'Событий в год на одного', 'Билетов на событие', 'Средняя цена, ₸', 'Оборот в год, ₸', 'Наш сбор, ₸']
for i, h in enumerate(heads):
    c = R.cell(4, 3 + i, h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True)
R.cell(4, 2, 'Сегмент').font = BOLD; R.cell(4, 2).fill = HEAD
segs = [
    ('Стендап и камеди-клубы', 60, 100, 80, 5000),
    ('Небольшие концерты и клубные вечера', 300, 24, 200, 7000),
    ('Независимые театры и труппы', 80, 60, 120, 5000),
    ('Платные конференции и митапы', 200, 6, 150, 15000),
    ('Локальные фестивали', 50, 2, 1500, 8000),
    ('Мастер-классы, лекции, экскурсии', 600, 24, 30, 4000),
]
for j, (name, n, e, t, p) in enumerate(segs):
    rr_ = 5 + j
    R.cell(rr_, 2, name)
    for i, v in enumerate((n, e, t, p)):
        c = R.cell(rr_, 3 + i, v); c.font = BLUE; c.number_format = TENGE
    R.cell(rr_, 7, f'=C{rr_}*D{rr_}*E{rr_}*F{rr_}').number_format = TENGE
    R.cell(rr_, 8, f'=G{rr_}*{names["fee"]}').number_format = TENGE
last = 5 + len(segs) - 1
R.cell(last + 1, 2, 'Итого (доступный рынок)').font = BOLD
R.cell(last + 1, 3, f'=SUM(C5:C{last})').number_format = TENGE
R.cell(last + 1, 7, f'=SUM(G5:G{last})').number_format = TENGE
R.cell(last + 1, 8, f'=SUM(H5:H{last})').number_format = TENGE
for col in (3, 7, 8): R.cell(last + 1, col).font = BOLD
t = last + 3
R.cell(t, 2, 'Доля рынка МСБ в 3-й год, базовый сценарий').font = BOLD
R.cell(t, 3, f'=IF(G{last+1}=0,0,SUM({yrange(2, "gmv", 3)})/G{last+1})').number_format = PCT
R.cell(t + 1, 2, 'Активных организаторов в конце 3-го года / всего в сегменте').font = BOLD
R.cell(t + 1, 3, f'=IF(C{last+1}=0,0,План!{cols[-1]}{block[2]["active"]}/C{last+1})').number_format = PCT
k = t + 3
R.cell(k, 2, 'Проверка порядка величин: Ticketon').font = H2
R.cell(k + 1, 2, 'Онлайн-заказов в месяц (2021)'); R.cell(k + 1, 3, 42810).font = BLUE; R.cell(k + 1, 3).number_format = TENGE
R.cell(k + 2, 2, 'Средний чек, $'); R.cell(k + 2, 3, 11.9).font = BLUE; R.cell(k + 2, 3).number_format = '0.0'
R.cell(k + 3, 2, 'Курс, ₸ за $'); R.cell(k + 3, 3, 480).font = BLUE; R.cell(k + 3, 3).number_format = TENGE
R.cell(k + 4, 2, 'Оборот онлайн-заказов в год, ₸').font = BOLD
R.cell(k + 4, 3, f'=C{k+1}*C{k+2}*C{k+3}*12').number_format = TENGE
R.cell(k + 5, 2, 'Доступный рынок МСБ к онлайн-обороту Ticketon, раз').font = BOLD
R.cell(k + 5, 3, f'=IF(C{k+4}=0,0,G{last+1}/C{k+4})').number_format = '0.0x'
R.cell(k + 5, 4, 'Больше 1: доступный рынок — все продажи сегмента, включая кассу, Instagram и Kaspi-переводы, а не только онлайн-билеты.').font = MUTED
R.cell(k + 6, 2, 'Наш оборот МСБ 3-го года (база), % от онлайн-оборота Ticketon 2021').font = BOLD
R.cell(k + 6, 3, f'=IF(C{k+4}=0,0,SUM({yrange(2, "gmv", 3)})/C{k+4})').number_format = PCT
R.cell(k + 1, 4, 'spec.md: оценка 2021 — выручка $6,1 млн, 42 810 онлайн-заказов в месяц, средний чек $11,9. Курс — допущение.').font = MUTED
g = k + 8
R.cell(g, 2, 'Крупный сегмент: проверка порядка величин').font = H2
R.cell(g + 1, 2, 'Крупных концертов и шоу от 2 000 мест в год по Казахстану')
R.cell(g + 1, 3, 450).font = BLUE; R.cell(g + 1, 3).number_format = TENGE
R.cell(g + 1, 4, ('Допущение: 400–500 в год. Ориентиры: Алматы — 334 концерта всех размеров в 2023 (zakon.kz), '
                  'плюс Дворец Республики в Алматы, площадки Астаны, арены и стадионы. Большинство — залы с ~2 000 проданных билетов. '
                  'Ошибка в 1,5 раза возможна.')).font = MUTED
R.cell(g + 2, 2, 'В том числе на аренах и стадионах, от 4 000 проданных билетов')
R.cell(g + 2, 3, 60).font = BLUE; R.cell(g + 2, 3).number_format = TENGE
R.cell(g + 2, 4, ('Допущение: Almaty Arena — около 24 в год (акимат Алматы, I полугодие 2024), Barys Arena — 10–15, '
                  'стадионы — 8–10, остальные — несколько. Наши крупные события по допущениям — от 4 000 билетов, '
                  'поэтому основной знаменатель — эта строка.')).font = MUTED
R.cell(g + 3, 2, 'Наших крупных событий в 3-й год (базовый)')
R.cell(g + 3, 3, f'={names["lev3_2"]}').font = GREEN; R.cell(g + 3, 3).number_format = NUM1
R.cell(g + 4, 2, 'Доля от событий арен и стадионов').font = BOLD
R.cell(g + 4, 3, f'=IF(C{g+2}=0,0,C{g+3}/C{g+2})').number_format = PCT
R.cell(g + 5, 2, 'Доля от всех событий от 2 000 мест').font = BOLD
R.cell(g + 5, 3, f'=IF(C{g+1}=0,0,C{g+3}/C{g+1})').number_format = PCT
for rr_ in range(g + 1, g + 6):
    R.cell(rr_, 2).alignment = Alignment(wrap_text=True, vertical='top')

# ---------------------------------------------------------------- Риски и источники
K = wb.create_sheet('Риски и источники')
K['A1'] = 'Риски модели и источники'; K['A1'].font = TITLE
K.column_dimensions['A'].width = 34; K.column_dimensions['B'].width = 110
risks = [
    ('Приём денег за организаторов',
     'Если оплата за билеты проходит через нас и потом выплачивается организатору, это может требовать регистрации '
     'платёжной организации в Нацбанке РК (Закон «О платежах и платёжных системах»). Альтернатива — сплит-платёж '
     'эквайера (деньги сразу делятся между организатором и нами); в MVP сплит-платежи исключены (CLAUDE.md). Проверить с юристом до запуска.'),
    ('Налоговая база',
     'Налог 4% считается только с нашего сбора при агентской схеме (договор с организатором, мы — агент). '
     'Без неё налоговой базой может считаться весь оборот билетов — модель тогда неверна.'),
    ('Ставка сбора',
     'При 5% базовый сценарий за три года не окупается: маржа с билета МСБ около 120 ₸, а эквайринг берётся со всей суммы '
     'оплаты и от ставки не зависит. Окупаемость за три года даёт ставка около 6% с крупным сегментом (33-й месяц) '
     'и около 6,5% без него (34-й месяц). Ставка — настройка продукта '
     '(SERVICE_FEE_BPS), её можно менять без кода; старые заказы сохраняют свой сбор.'),
    ('Численности рынка',
     'Число организаторов по сегментам и число крупных событий — допущения. Интервью до защиты не проводятся (решение владельца '
     '2026-10-09); численности проверяются данными первых месяцев после запуска.'),
    ('Крупный сегмент',
     'Числа без интервью. Билетные операторы связаны эксклюзивами и предпродажей в банковских приложениях '
     '(Ticketon — Freedom SuperApp, kino.kz — Halyk); крупные промоутеры торгуются об условиях. 5% с покупателя и 0% с организатора — наша позиция, а не их согласие: если промоутер '
     'потребует половину сбора, маржа базового события падает примерно с 2,3 до 0,1 млн ₸, а после затрат на продажу и сопровождение событие уходит '
     'в минус около 0,9 млн ₸. Плановое расширение подов перед стартом '
     'продаж не реализовано (docs/adr/032-kubernetes.md, «Открыто»), реактивный HPA на короткий шторм опаздывает, поэтому мощность '
     'на первые крупные события поднимаем вручную заранее. Авансы промоутерам потребуют статуса платёжной организации или сплит-платежа '
     'эквайера (риск «Приём денег за организаторов»). Отмена события отдельно не учтена: сбор возвращается, затраты остаются.'),
    ('Конкуренты', 'Ticketon и kino.kz могут снизить сбор для МСБ; наш ответ — бесплатный вход и своя схема зала.'),
]
K['A3'] = 'Риск'; K['B3'] = 'Описание и что делать'
for c in ('A3', 'B3'): K[c].font = BOLD; K[c].fill = HEAD
for j, (a, b) in enumerate(risks):
    K.cell(4 + j, 1, a).font = BOLD
    c = K.cell(4 + j, 2, b); c.alignment = Alignment(wrap_text=True, vertical='top')
    K.row_dimensions[4 + j].height = 48 if len(b) < 400 else 96
s0 = 4 + len(risks) + 2
K.cell(s0, 1, 'Источники').font = H2
srcs = [
    ('Сервисный сбор ~10%, «10/10»', 'https://support.ticketscloud.com/ (инструкция для организаторов, условия сделок)'),
    ('Эквайринг Kaspi Pay', 'https://guide.kaspi.kz/partner/ru/account/opening/q258'),
    ('Эквайринг Halyk ePay', 'https://halykbank.kz/en/business/payment/epay'),
    ('Сравнение эквайринга, 2026', 'https://finkaz.kz/acquiring'),
    ('Упрощённый режим 4% с 2026', 'https://digitalbusiness.kz/2025-08-07/uznali-kogda-objyavyat-stavki-po-novoy-uproshchenke-na-2026-god/'),
    ('Ticketon, ориентир 2021', 'docs/spec.md, раздел «Бизнес-модель»'),
    ('Крупный сегмент (поиск 2026-10-09: числа сверены по выдаче поиска двумя независимыми проверками, сами страницы из среды разработки не открывались)', ''),
    ('Almaty Arena, вместимость', 'https://en.wikipedia.org/wiki/Almaty_Arena'),
    ('Almaty Arena и Дворец Республики, I полугодие 2024', 'https://www.gov.kz/memleket/entities/almaty/press/news/details/800411?lang=ru'),
    ('Алматы, 334 концерта в 2023', 'https://www.zakon.kz/press-relizy/6420112-kulturnye-sobytiya-goda-v-almaty-334-kontserta-bolee-16-mln-posetiteley-i-zvezdy-mirovogo-urovnya.html'),
    ('Barys Arena, вместимость', 'https://newtimes.kz/kultura-i-sport/18200-v-astane-sostoyalos-otkrytie-mnogofunktsionalnogo-ledovogo-dvortsa'),
    ('Astana Arena, вместимость', 'https://en.wikipedia.org/wiki/Astana_Arena'),
    ('Центральный стадион Алматы, вместимость', 'https://en.wikipedia.org/wiki/Almaty_Central_Stadium'),
    ('Би-2 в Almaty Arena, ~8 000 зрителей, 2022', 'https://informburo.kz/novosti/kak-proshyol-koncert-bi-2-v-almaty'),
    ('Kazakh Tourism: 5 крупных концертов в 2026', 'https://bizmedia.kz/2026-01-27-mirovye-zvyozdy-snova-priedut-v-kazahstan-timbaland-busta-rhymes-i-nikol-sherzinger/'),
    ('Средний чек Ticketon на концерты, январь–июль 2024', 'https://bes.media/news/kontserti-futbolnie-matchi-ili-kino-kuda-bolshe-vsego-pokupayut-bileti-kazahstantsi/'),
    ('Ticketon, итоги 2024', 'https://finratings.kz/news/1243-prodazhi-biletov-na-platforme-freedom-ticketon-udvoilis-v-2024-godu/'),
    ('Макс Корж, Центральный стадион, 2025', 'https://digitalbusiness.kz/2025-03-20/prodali-polovinu-stadiona-za-paru-chasov-chto-izvestno-o-kontserte-maksa-korzha-v-almati/'),
    ('Skillet, Алматы, 2024', 'https://freedom-travel.kz/media/news/skillet-astana-and-almaty/'),
    ('Backstreet Boys, цены 2025', 'https://bes.media/news/kontserti-backstreet-boys-v-astane-i-almati-bileti-budut-stoit-ot-40-do-200-tisyach-tenge/'),
    ('kino.kz (Halyk) и Backstreet Boys', 'https://bes.media/news/kinokz-ot-halyk-2-mln-prosmotrov-i-21-tis-biletov-na-backstreet-boys-za-sutki/'),
    ('Канье Уэст: цены и сбор 10%', 'https://rus.azattyq-ruhy.kz/ar-trend/112114-stali-izvestny-tseny-na-bilety-na-pervyi-kontsert-kane-uesta-v-almaty'),
    ('Энрике Иглесиас, цены 2026', 'https://forbes.kz/articles/enrike-iglesias-v-almaty-i-astane-skolko-stoyat-bilety-95372c'),
    ('Дженнифер Лопес: сбой Ticketon и фиксированная комиссия', 'https://kapital.kz/business/136156/koncert-dzhennifer-lopes-timur-turlov-rasskazal-o-poteryah-freedom-ticketon-iz-za-sboya.html'),
    ('Сбой Ticketon на матче «Кайрат» — «Реал»', 'https://finratings.kz/news/7344-freedom-ticketon-upal-pri-prodazhe-biletov-na-match-kairat-real/amp'),
    ('Ticketon: билеты на несуществующие места', 'https://finratings.kz/news/2847-ticketon-prodal-bilety-na-nesushchestvuiushchie-mesta/'),
    ('Ticketon расторг партнёрство с Яндекс Афишей, 2024', 'https://forbes.kz/articles/freedom-ticketon-rastorg-partnerstvo-syandeks-afisha'),
]
for j, (a, b) in enumerate(srcs):
    K.cell(s0 + 1 + j, 1, a).font = BOLD if not b else BLACK; K.cell(s0 + 1 + j, 2, b)

for ws in wb.worksheets:
    style_all(ws)
    # Печать: альбомная, по ширине страницы.
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False if ws.title in ('Итоги', 'Юнит-экономика', 'Сравнение моделей', 'Риски и источники') else True
wb.save(OUT)
print('saved', OUT)
